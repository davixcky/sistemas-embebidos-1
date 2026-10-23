"""Genera PDF, PowerPoint editable y guion a partir de evidencia del proyecto."""
from pathlib import Path
import json
from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'docs/presentacion/archivo';OUT.mkdir(parents=True,exist_ok=True)
FONT='/usr/share/fonts/liberation/LiberationSans-Regular.ttf';BOLD='/usr/share/fonts/liberation/LiberationSans-Bold.ttf'
pdfmetrics.registerFont(TTFont('Body',FONT));pdfmetrics.registerFont(TTFont('Bold',BOLD))
S=[
('Medir el total y un subcircuito', 'Medidor de energía con ESP32 · Avance 2', '3d-medidor-cerrado.png', ['120 V · 60 Hz','Dos pinzas de corriente y un canal de tensión','David Orozco · Jose Torres · Leo Rojas'], 'David · 45 s','La factura muestra un total mensual. Este proyecto busca distinguir el consumo total del de un circuito, por ejemplo el aire acondicionado. La imagen es el diseño CAD, no una fotografía del prototipo. El avance se concentra en el acondicionamiento, la placa y el gabinete.'),
('Tres objetivos verificables', 'Diseñar, construir y validar el medidor', None, ['Diseñar las entradas para 60 A en CH1 y 30 A en CH2, dentro del rango útil del ADC.','Construir e integrar el prototipo y superar las pruebas eléctricas y de resistencia.','Programar y validar potencia con error de ±3 % o menor entre 1 A y 30 A.'], 'David · 55 s', 'Objetivo general: desarrollar un medidor monofásico no invasivo de dos canales, con tensión aislada, cálculo eléctrico y publicación por Wi-Fi. El error del 3 % es una meta; no se ha demostrado en el equipo físico. El primer avance proponía tres pinzas: este corte implementa dos y conserva entradas de reserva.'),
('La arquitectura conserva el cálculo en el ESP32', 'Pinzas y tensión → acondicionamiento → ADC1 → cálculo → pantalla y web', 'diagrama-de-bloques.png', [], 'David · 60 s', 'La corriente se mide abrazando solo la fase. El ZMPT101B aporta la señal de tensión aislada. Se calcula Vrms, Irms y potencia activa a partir de las muestras; la energía integra potencia en el tiempo. ADC1 permite adquirir mientras funciona Wi-Fi. Mostrar en el diagrama la ruta completa de las señales.'),
('La diferencia es poder explicar la medición', 'Referencias del diseño y fundamento del cálculo', None, ['Shelly EM: medición comercial de dos canales.','IoTaWatt y OpenEnergyMonitor: referencias abiertas para sensores y adquisición.','Este diseño: resistencias de carga, ADC del ESP32 y cálculo propio.','P = promedio de v·i; S = Vrms·Irms; FP = P/S.'], 'David · 60 s', 'No se presenta como un instrumento más preciso que los equipos comerciales. El valor académico es conocer y justificar cada etapa. Las fuentes están en las referencias 10 a 14 del README. Q calculada como raíz de S²−P² requiere cuidado con cargas no sinusoidales: no separa por sí sola toda la distorsión.'),
('Los materiales corresponden a cada etapa', 'Lista completa y referencias: sección 6 del informe', None, ['Adquisición: ESP32-DevKitC V4, dos SCT-013-000 y un ZMPT101B.','Entradas: 22 Ω / 47 Ω, polarización de 1,32 V, RC y BAT85.','Interfaz: OLED SSD1306, dos pulsadores y LED.','Montaje: placa perforada, conectores, gabinete y fuente aislada de 5,1 V.'], 'Jose · 50 s', 'Indicar únicamente las piezas que realmente están compradas o montadas. La lista del informe es la lista prevista, no un inventario de compra. El banco tiene una lista propia. No llamar PCB fabricada a la placa perforada: la equivalencia con el requisito del curso queda por confirmar con el docente.'),
('La resistencia de carga fija la señal de corriente', 'Diseño nominal a 27 °C · sin sobrecorriente', None, ['CH1: 60 A → 30 mA secundarios → 22 Ω → 0,66 Vrms.','CH2: 30 A → 15 mA secundarios → 47 Ω → 0,705 Vrms.','Polarización: 3,3 × 10 / (15 + 10) = 1,32 V.','Filtro: 3,3 kΩ + 47 nF; corte cercano a 1 kHz.'], 'Jose · 70 s', 'La señal alterna se centra en 1,32 V para evitar voltajes negativos en funcionamiento normal. El burden se suelda en el jack. Eso no protege una pinza desenchufada; no se desconecta bajo carga. Las protecciones BAT85 necesitan revisión: no basta con que el modelo nominal quede cerca del máximo absoluto.'),
('La distribución coincide con el esquemático', 'Placa perforada de 150 × 90 mm', 'placa-perforada-componentes.png', ['ERC: 0 violaciones','DRC guardado: 0 conexiones pendientes y 9 solapes de patios'], 'Jose · 55 s', 'Distinguir conectividad de montaje. Los nueve solapes de patios están documentados; no son nueve cables faltantes. Queda revisar el ajuste físico. La placa no tiene orificios libres para los soportes M3 del gabinete: resolver clips o borde de montaje antes de fabricar. Abrir el PDF del esquemático si se pide seguir una conexión.'),
('El gabinete separa la red de la electrónica', '200 × 150 × 75 mm · modelo CAD', '3d-medidor-abierto.png', ['Compartimento para ZMPT101B y fusible','OLED, pulsadores y LED en la tapa','Video: ../video/medidor-animado.mp4'], 'Jose · 65 s', 'Reproducir el video corto y explicar la separación de zonas. El movimiento de la tapa es un despiece virtual. El gabinete CAD no acredita fabricación, aislamiento ni resistencia mecánica. El cableado del panel y la fijación deben verificarse en el montaje real.'),
('El banco permite comparar una carga conocida', 'Protocolo de verificación · no es un resultado experimental', '3d-banco-abierto.png', ['Una fase por pinza','Referencia: pinza amperimétrica o PZEM','Calibración y prueba de 2 h pendientes'], 'Leo · 55 s', 'Explicar por qué una pinza alrededor de fase y neutro juntos no mide el consumo. La verificación de continuidad y aislamiento antecede a la energización. El banco está diseñado; no afirmar que se construyó o calibró sin evidencia. Las pruebas de resistencia figuran pendientes en la tabla 8 del informe.'),
('Dos simuladores coinciden; la protección necesita margen', 'Xyce 7.10 y ngspice 46 · 24 casos por simulador', None, ['Nominal: CH1 0,392–2,246 V; CH2 0,329–2,309 V.','Diferencia de voltajes entre solvers: menor de 10 µV.','CH2, 100 A a 15 °C: −0,334 a 3,627 V; fuera del límite usado.','Adaptador de 4,75 V: 4,306 V antes del regulador; no cumple el criterio de 4,4 V.'], 'Leo · 70 s', 'Los dos solvers usan los mismos modelos: su coincidencia verifica el cálculo numérico, no la exactitud física. El barrido térmico es una prueba de sensibilidad con el modelo ajustado de diodo, no un modelo térmico validado por fabricante. No se modelan saturación de la pinza, ADC real ni disparo del PTC. La protección no está cerrada.'),
('El software ya admite una demostración con datos simulados', 'Adelanto del corte 3 · no sustituye el prototipo de C2', 'ui-tablero-escritorio.png', ['10 pruebas web aprobadas; compilación de producción correcta','Firmware compilado y cálculo probado con señales sintéticas'], 'Leo · 55 s', 'La captura proviene del simulador, no del nodo físico. Se puede mostrar la recepción, historial, energía y alertas. El LSTM usa datos públicos y no demuestra precisión en una vivienda local. No dedicar más tiempo al software que al hardware exigido en este corte.'),
('Lo pendiente está identificado', 'Estado documental: 23 de octubre de 2026', None, ['Diseño, CAD, simulaciones y pruebas de software disponibles.','Montaje, gabinete físico y HardExplosion: sin evidencia de cierre en el repositorio.','Antes de montar: revisar clamp, alimentación y fijación de la placa.','Después: calibrar, conectar el nodo real y comprobar la meta de ±3 %.'], 'Leo · 55 s', 'La metodología tiene tres etapas: diseño, implementación física y validación. La compra presentó retrasos y el montaje está registrado en curso desde el 16 de octubre. Actualizar solo si el grupo aporta evidencia real. La entrega documental debe quedar en GitHub/GitLab; este directorio está sin seguimiento y no se confirmó su publicación. Cerrar con los próximos pasos concretos, sin declarar terminado el prototipo.'),
('Anexo A · Esquemático completo', 'Abrir hardware/medidor-energia.pdf para ampliar conexiones', 'esquematico.png', [], 'Consulta', 'Referencia ampliable: hardware/medidor-energia.pdf. El ERC comprueba reglas eléctricas; no certifica precisión ni seguridad del montaje.'),
('Anexo B · Tensión y alimentación', 'Valores y límites que se deben sustentar', None, ['ZMPT101B: salida nominal ajustada a 2,5 V ± 0,65 V.','Divisor: 11 kΩ / 13 kΩ; relación 0,542.','Filtro de tensión: 27 nF y resistencia equivalente 5,96 kΩ.','Fuente: 5,1 V aislada; comprobar caídas reales en PTC y diodo.'], 'Consulta', 'El filtro de tensión tiene una constante de tiempo cercana a la de corriente. D1 bloquea inversión de la fuente y retorno hacia el adaptador; no justifica por sí solo conectar simultáneamente USB y fuente. Espressif indica fuentes mutuamente excluyentes.'),
('Anexo C · Evidencias y fuentes', 'Archivos locales para responder preguntas', None, ['Informe completo: README.md; referencias en la sección 9.','Esquemáticos y verificaciones: hardware/.','Comparación de simuladores: simulacion/xyce/resultados.txt.','Rúbrica: material del curso, sección 1.1.9.3.'], 'Consulta', 'Fuente de la rúbrica: https://nemog3762.gitlab.io/material/compiled/Sistemas%20embebidos%201.html. Figuras realizadas con KiCad, FreeCAD y scripts de dibujo, no con generación de imágenes por IA. Consultar docs/presentacion/control-entrega.md antes de salir.')]
(OUT/'contenido.json').write_text(json.dumps(S,ensure_ascii=False,indent=2))
W,H=960,540
c=canvas.Canvas(str(OUT/'presentacion-c2.pdf'),pagesize=(W,H));c.setTitle('Medidor de energía · Avance 2');c.setAuthor('David Orozco · Jose Torres · Leo Rojas')
p=Presentation();p.slide_width=Inches(13.3333);p.slide_height=Inches(7.5)
INK='#172B35';MUTED='#50636D';ACC='#007D82'
def wrap(t,font,size,width):
 lines=[];line=''
 for word in t.split():
  trial=(line+' '+word).strip()
  if pdfmetrics.stringWidth(trial,font,size)>width and line:lines.append(line);line=word
  else:line=trial
 if line:lines.append(line)
 return lines

def text(sl,t,x,y,w,size=22,bold=False,color=INK):
 lines=wrap(t,'Bold' if bold else 'Body',size,w)
 c.setFillColor(color);c.setFont('Bold' if bold else 'Body',size)
 for j,line in enumerate(lines):c.drawString(x,H-y-size-j*size*1.22,line)
 box=sl.shapes.add_textbox(Inches(x/72), Inches(y/72), Inches(w/72), Inches((len(lines)*size*1.28+8)/72))
 tf=box.text_frame;tf.word_wrap=False;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for j,line in enumerate(lines):
  par=tf.paragraphs[0] if j==0 else tf.add_paragraph();par.text=line
  par.font.name='Liberation Sans';par.font.size=Pt(size);par.font.bold=bold;par.font.color.rgb=RGBColor.from_string(color[1:])
 return len(lines)*size*1.22

def pic(sl,path,x,y,w,h):
 iw,ih=Image.open(path).size;scale=min(w/iw,h/ih);ww,hh=iw*scale,ih*scale;xx=x+(w-ww)/2;yy=y+(h-hh)/2
 c.drawImage(str(path),xx,H-yy-hh,ww,hh)
 sl.shapes.add_picture(str(path), Inches(xx/72), Inches(yy/72),width=Inches(ww/72),height=Inches(hh/72))
for i,(title,sub,img,bullets,who,note) in enumerate(S,1):
 sl=p.slides.add_slide(p.slide_layouts[6]);c.setFillColor('#FAFBFC');c.rect(0,0,W,H,fill=1,stroke=0)
 sl.background.fill.solid();sl.background.fill.fore_color.rgb=RGBColor.from_string('FAFBFC')
 text(sl,f'C2 / ABPr     SISTEMAS EMBEBIDOS 1',40,22,800,10,True,ACC)
 text(sl,f'{i:02d}. {title}',40,53,880,28,True)
 text(sl,sub,40,99,880,15,False,MUTED)
 if img:
  if bullets:
   pic(sl,ROOT/'img'/img,365,143,555,320)
   video=ROOT/'docs/video/medidor-animado.mp4'
   if i==8 and video.exists():
    sl.shapes.add_movie(str(video), Inches(365/72), Inches(147/72), Inches(555/72), Inches(312.1875/72), poster_frame_image=str(ROOT/'docs/video/portada.jpg'), mime_type='video/mp4')
   y=166
   for b in bullets:y+=text(sl,b,40,y,290,21)+25
  else:pic(sl,ROOT/'img'/img,40,137,880,344)
 else:
  y=159
  for b in bullets:
   y+=text(sl,b,55,y,840,25 if len(bullets)<4 else 23)+25
 text(sl,'CUC · David Orozco / Jose Torres / Leo Rojas',40,507,700,10,False,MUTED)
 text(sl,f'{i:02d} / {len(S):02d}',860,507,80,10,False,MUTED)
 sl.notes_slide.notes_text_frame.text=f'{who}\n\n{note}'
 c.showPage()
c.save();p.save(OUT/'presentacion-c2.pptx')
# A readable rehearsal script; not prose to read verbatim during the talk.
c=canvas.Canvas(str(OUT/'guion-exposicion.pdf'),pagesize=(595,842));c.setTitle('Guion de exposición C2')
y=790
for i,(title,sub,img,bullets,who,note) in enumerate(S,1):
 lines=wrap(note,'Body',11,495);titlelines=wrap(f'{i:02d}. {title}','Bold',14,495);need=60+len(lines)*15+len(titlelines)*18
 if y-need<60:c.showPage();y=790
 c.setFillColor(INK);c.setFont('Bold',14)
 for line in titlelines:c.drawString(50,y,line);y-=18
 y-=5
 c.setFont('Body',11);c.setFillColor(ACC);c.drawString(50,y,who);y-=24;c.setFillColor(INK)
 for line in lines:c.drawString(50,y,line);y-=15
 y-=26
c.save()
print('PDF, PPTX y guion generados:',len(S),'páginas; 12 principales y 3 anexos')

# Checklist printable, including rubric mapping and evidence tables.
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
import html,re
styles=getSampleStyleSheet()
for key in ['BodyText','Heading1','Heading2']:
 styles[key].fontName='Bold' if key.startswith('Heading') else 'Body'
styles['BodyText'].fontSize=9;styles['BodyText'].leading=12
styles['Heading1'].fontSize=17;styles['Heading2'].fontSize=12
parts=[];table=[]
def clean(line):
 return html.escape(re.sub(r'\[([^\]]+)\]\([^)]+\)',r'\1',line).replace('- [ ]','[  ]').replace('**',''))
def flush_table():
 if not table:return
 rows=[[Paragraph(clean(cell),styles['BodyText']) for cell in row] for row in table]
 widths=[255,244] if len(table[0])==2 else [155,180,164]
 t=Table(rows,colWidths=widths,repeatRows=1,hAlign='LEFT')
 t.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E8EFF1')),('LINEBELOW',(0,0),(-1,-1),.3,colors.HexColor('#CED8DC')),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
 parts.append(t);table.clear()
for line in (OUT/'control-entrega.md').read_text().splitlines():
 if line.startswith('|'):
  cells=[x.strip() for x in line.strip('|').split('|')]
  if not all(re.fullmatch('[: -]+',x) for x in cells):table.append(cells)
  continue
 flush_table()
 if not line.strip():parts.append(Spacer(1,6));continue
 style='Heading1' if line.startswith('# ') else 'Heading2' if line.startswith('## ') else 'BodyText'
 parts.append(Paragraph(clean(re.sub(r'^#{1,2} ','',line)),styles[style]))
flush_table()
SimpleDocTemplate(str(OUT/'control-entrega.pdf'),title='Control de entrega C2',author='Equipo C2 ABPr',leftMargin=48,rightMargin=48,topMargin=42,bottomMargin=42).build(parts)
